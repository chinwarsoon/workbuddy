# -*- coding: utf-8 -*-
"""Analyse TWRP C3B2 MDL + Submittal Log for monthly submission dashboard."""
import openpyxl, collections, datetime, json, os

BASE = r"C:/Users/frank/WorkBuddy/workbuddy/MDL"
MDL = os.path.join(BASE, "TWRP C3B2 - Master Document List.xlsx")
LOG = os.path.join(BASE, "TWRP C3B2 - Document Submittal Log.xlsx")

def ym(v):
    if isinstance(v, datetime.datetime):
        return v.strftime("%Y-%m")
    if isinstance(v, datetime.date):
        return v.strftime("%Y-%m")
    return None

def clean(v):
    if v is None:
        return None
    if isinstance(v, str):
        s = v.strip()
        if s == "" or s.startswith("#"):
            return None
        return s
    return v

# ---------------- MDL ----------------
wb = openpyxl.load_workbook(MDL, read_only=True, data_only=True)
skel = ["Project", "Construction", "QAQC", "ICA", "Process", "BS", "Mech",
        "Piping", "Electrical", "CSA", "BIM"]
mdl_rows = []
disc_stat = collections.Counter()
status_vals = collections.Counter()
appr_vals = collections.Counter()
planned_ym = collections.Counter()
fore_ym = collections.Counter()
actual_ym = collections.Counter()
stage_c = collections.Counter()
doctype_c = collections.Counter()
for sh in skel:
    ws = wb[sh]
    it = ws.iter_rows(min_row=2, values_only=True)
    for r in it:
        sno = clean(r[0])
        stage = clean(r[1])
        disc = clean(r[2])
        prefix = clean(r[7])
        doctype = clean(r[8])
        dcode = clean(r[9])
        docno = clean(r[12])
        title = clean(r[13])
        vendor = clean(r[17])
        planned = r[18]
        fore = r[19]
        actual = r[20]
        sub_stat = clean(r[22])
        appr = clean(r[24])
        if docno is None and title is None and stage is None:
            continue
        if sno is None and stage is None:
            continue  # group header
        disc_stat[sh] += 1
        stage_c[stage] += 1
        doctype_c[doctype] += 1
        status_vals[sub_stat] += 1
        appr_vals[appr] += 1
        p, f, a = ym(planned), ym(fore), ym(actual)
        if p: planned_ym[p] += 1
        if f: fore_ym[f] += 1
        if a: actual_ym[a] += 1
        mdl_rows.append(dict(disc=sh, stage=stage, prefix=prefix, doctype=doctype,
                             docno=docno, title=title, vendor=vendor,
                             planned=p, fore=f, actual=a, sub_stat=sub_stat, appr=appr))

print("=" * 70)
print("MDL: %d document rows" % len(mdl_rows))
print("by discipline:", dict(disc_stat))
print("by stage:", dict(stage_c.most_common()))
print("by doc type:", dict(doctype_c.most_common(12)))
print("Submission Status values:", dict(status_vals))
print("Latest Approval Status values:", dict(appr_vals))
print("planned month range:", min(planned_ym) if planned_ym else None, "~", max(planned_ym) if planned_ym else None, "months:", len(planned_ym))
print("forecast month range:", min(fore_ym) if fore_ym else None, "~", max(fore_ym) if fore_ym else None, "months:", len(fore_ym))
print("actual month range:", min(actual_ym) if actual_ym else None, "~", max(actual_ym) if actual_ym else None, "months:", len(actual_ym))
print()
print("PLANNED by month:")
for k in sorted(planned_ym):
    print("   ", k, planned_ym[k])
print()
print("FORECAST by month:")
for k in sorted(fore_ym):
    print("   ", k, fore_ym[k])
print()
print("ACTUAL ISSUE by month:")
for k in sorted(actual_ym):
    print("   ", k, actual_ym[k])

# ---------------- Submittal Log ----------------
wb2 = openpyxl.load_workbook(LOG, read_only=True, data_only=True)
ws = wb2["SubmittalLog"]
log_rows = []
cur_disc = None
grp = []
for r in ws.iter_rows(min_row=2, values_only=True):
    sno = clean(r[0]); pno = clean(r[1]); prev = clean(r[2]); subj = clean(r[3])
    if sno is None and subj and subj in ("ICA", "Electrical", "Mechanical", "Building Services", "Process", "CSA", "Piping", "BIM", "Project", "Construction", "QAQC"):
        cur_disc = subj
        continue
    if sno is None:
        continue
    docno = clean(r[4]); rev = clean(r[5]); by = clean(r[6])
    sdate = r[7]; rperiod = clean(r[8]); target = r[9]; aret = r[10]
    turn = clean(r[11]); delay = clean(r[12])
    sstat = clean(r[13]); ast = clean(r[14]); resub = clean(r[15])
    log_rows.append(dict(sno=sno, pno=pno, disc=cur_disc, subj=subj, docno=docno,
                         rev=rev, by=by, sdate=sdate, rperiod=rperiod, target=target,
                         aret=aret, turn=turn, delay=delay, sstat=sstat, ast=ast, resub=resub))

print()
print("=" * 70)
print("SUBMITTAL LOG: %d rows" % len(log_rows))
sm = collections.Counter(ym(r["sdate"]) for r in log_rows if ym(r["sdate"]))
print("submission by month:")
for k in sorted(sm):
    print("   ", k, sm[k])
print("by discipline:", dict(collections.Counter(r["disc"] for r in log_rows)))
print("submission status:", dict(collections.Counter(r["sstat"] for r in log_rows)))
print("approval status:", dict(collections.Counter(r["ast"] for r in log_rows)))
print("resubmission:", dict(collections.Counter(r["resub"] for r in log_rows)))
print("submitted by:", dict(collections.Counter(r["by"] for r in log_rows).most_common()))
print("review period:", dict(collections.Counter(r["rperiod"] for r in log_rows)))
turns = [float(r["turn"]) for r in log_rows if r["turn"] not in (None, "")]
delays = [float(r["delay"]) for r in log_rows if r["delay"] not in (None, "")]
print("turnaround: n=%d min=%.0f max=%.0f avg=%.1f" % (len(turns), min(turns), max(turns), sum(turns)/len(turns)))
print("delay: n=%d min=%.0f max=%.0f avg=%.1f" % (len(delays), min(delays), max(delays), sum(delays)/len(delays)))
print("returned (actual return date present):", sum(1 for r in log_rows if r["aret"]))
print("docno sample:", [r["docno"] for r in log_rows[:5]])
